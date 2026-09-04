"""Importación / exportación de tareas vía Excel (.xlsx), en dos archivos
separados que reflejan el flujo real de trabajo:

  1) Cronograma planificado — crea las tareas y su cronograma teórico
     (se carga antes de arrancar la obra, o para sumar tareas nuevas).
       columnas: nombre | tag | descripcion | inicio_plan | fin_plan

  2) Tiempos reales — carga el inicio/fin real de cada tarea y sus
     demoras, todo en un solo archivo. Cada tarea se identifica por su
     TAG (debe existir ya en la obra). Si una tarea tuvo más de una
     demora, se repite su fila (mismo tag) una vez por demora.
       columnas: tag | inicio_real | fin_real | demora_inicio | demora_fin | motivo_demora

Las columnas de fecha/hora aceptan tanto celdas de fecha de Excel como
texto en formato "AAAA-MM-DD HH:MM" (se normalizan al formato interno
"AAAA-MM-DDTHH:MM" que usa toda la aplicación).
"""
import io
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

CRONOGRAMA_HEADERS = ["nombre", "tag", "descripcion", "inicio_plan", "fin_plan"]
REALES_HEADERS = ["tag", "inicio_real", "fin_real", "demora_inicio", "demora_fin", "motivo_demora"]

INTERNAL_FMT = "%Y-%m-%dT%H:%M"


def _hoja_con_estilo(wb, titulo, headers):
    ws = wb.active
    ws.title = titulo
    header_fill = PatternFill("solid", fgColor="1C1C1C")
    header_font = Font(color="FFFFFF", bold=True)
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        ws.column_dimensions[get_column_letter(col)].width = max(16, len(headers[col - 1]) + 4)
    return ws


def _hoja_instrucciones(wb, lineas):
    ws = wb.create_sheet("Instrucciones")
    for i, line in enumerate(lineas, start=1):
        ws.cell(row=i, column=1, value=line)
    ws.column_dimensions["A"].width = 100
    ws["A1"].font = Font(bold=True, size=13)


def generar_plantilla_cronograma(items_existentes=None):
    """Workbook para crear tareas + cronograma planificado."""
    wb = Workbook()
    ws = _hoja_con_estilo(wb, "Cronograma", CRONOGRAMA_HEADERS)
    if items_existentes:
        for it in items_existentes:
            ws.append([
                it["nombre"], it["tag_equipo"] or "", it["descripcion"] or "",
                it["fecha_inicio_plan"] or "", it["fecha_fin_plan"] or "",
            ])
    else:
        ws.append(["Bomba centrífuga P-101", "P-101", "Desmontaje, limpieza y alineación.",
                    "2026-09-01 08:00", "2026-09-05 17:00"])

    _hoja_instrucciones(wb, [
        "Cómo usar esta plantilla — Cronograma planificado",
        "",
        "Una fila por cada ítem / equipo a intervenir en la obra.",
        "  - nombre: obligatorio.",
        "  - tag: código del equipo. Se usa para identificar la tarea en esta y en la plantilla",
        "    de 'Tiempos reales'. Si volvés a importar un tag que ya existe en la obra, esa tarea",
        "    se actualiza en vez de duplicarse.",
        "  - inicio_plan / fin_plan: fecha y hora planificadas.",
        "  - Las fechas podés escribirlas como fecha de Excel o como texto 'AAAA-MM-DD HH:MM'.",
        "",
        "Podés borrar la fila de ejemplo antes de cargar tus datos.",
    ])
    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def generar_plantilla_reales(items_existentes=None, demoras_por_item=None):
    """Workbook para cargar inicio/fin real + demoras de tareas ya creadas."""
    wb = Workbook()
    ws = _hoja_con_estilo(wb, "TiemposReales", REALES_HEADERS)
    demoras_por_item = demoras_por_item or {}
    if items_existentes:
        for it in items_existentes:
            tag = it["tag_equipo"] or ""
            demoras = demoras_por_item.get(it["id"], [])
            if demoras:
                first = demoras[0]
                ws.append([tag, it["fecha_inicio"] or "", it["fecha_fin"] or "",
                           first["fecha_inicio"], first["fecha_fin"], first["motivo"] or ""])
                for dem in demoras[1:]:
                    ws.append([tag, "", "", dem["fecha_inicio"], dem["fecha_fin"], dem["motivo"] or ""])
            else:
                ws.append([tag, it["fecha_inicio"] or "", it["fecha_fin"] or "", "", "", ""])
    else:
        ws.append(["P-101", "2026-09-01 08:30", "2026-09-05 15:30",
                    "2026-09-02 12:00", "2026-09-02 13:30", "Espera de repuesto en depósito."])

    _hoja_instrucciones(wb, [
        "Cómo usar esta plantilla — Tiempos reales y demoras",
        "",
        "Una fila por tarea (identificada por su TAG, que ya tiene que existir en la obra —",
        "primero importá o cargá el cronograma planificado).",
        "  - tag: obligatorio, tiene que coincidir con el tag de una tarea de la obra.",
        "  - inicio_real / fin_real: fecha y hora real de inicio/fin de la tarea (dejar vacío si",
        "    todavía no arrancó / no terminó).",
        "  - demora_inicio / demora_fin / motivo_demora: opcional. Si la tarea tuvo una demora,",
        "    completá estas tres columnas.",
        "",
        "¿La tarea tuvo más de una demora? Repetí una fila extra con el mismo tag: dejá",
        "inicio_real y fin_real vacíos en esa fila extra, y completá solo las columnas de la",
        "demora correspondiente.",
        "",
        "Las fechas podés escribirlas como fecha de Excel o como texto 'AAAA-MM-DD HH:MM'.",
    ])
    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _norm_fecha(value):
    """Normaliza un valor de celda (datetime de Excel o texto) al formato interno."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.strftime(INTERNAL_FMT)
    text = str(value).strip()
    if not text:
        return None
    for fmt in (INTERNAL_FMT, "%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).strftime(INTERNAL_FMT)
        except ValueError:
            continue
    return None  # valor no reconocible: se descarta en vez de romper la importación


def _hoja_principal(wb):
    """Devuelve la primera hoja de datos (ignora una posible hoja 'Instrucciones')."""
    for name in wb.sheetnames:
        if name != "Instrucciones":
            return wb[name]
    return wb.active


def leer_cronograma(file_storage):
    """Lee el Excel de cronograma planificado. Devuelve (tareas, errores)."""
    wb = load_workbook(file_storage, data_only=True)
    ws = _hoja_principal(wb)
    errores = []
    tareas = []
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    for idx, row in enumerate(rows, start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        row = list(row) + [None] * (len(CRONOGRAMA_HEADERS) - len(row))
        nombre = (str(row[0]).strip() if row[0] else "")
        if not nombre:
            errores.append(f"Fila {idx}: sin nombre, se saltea.")
            continue
        tareas.append({
            "nombre": nombre,
            "tag": (str(row[1]).strip() if row[1] else ""),
            "descripcion": (str(row[2]).strip() if row[2] else ""),
            "fecha_inicio_plan": _norm_fecha(row[3]),
            "fecha_fin_plan": _norm_fecha(row[4]),
        })
    return tareas, errores


def leer_reales(file_storage):
    """Lee el Excel de tiempos reales + demoras. Devuelve (filas, errores)."""
    wb = load_workbook(file_storage, data_only=True)
    ws = _hoja_principal(wb)
    errores = []
    filas = []
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    for idx, row in enumerate(rows, start=2):
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        row = list(row) + [None] * (len(REALES_HEADERS) - len(row))
        tag = (str(row[0]).strip() if row[0] else "")
        if not tag:
            errores.append(f"Fila {idx}: sin tag, se saltea.")
            continue
        filas.append({
            "tag": tag,
            "fecha_inicio": _norm_fecha(row[1]),
            "fecha_fin": _norm_fecha(row[2]),
            "demora_inicio": _norm_fecha(row[3]),
            "demora_fin": _norm_fecha(row[4]),
            "motivo_demora": (str(row[5]).strip() if row[5] else ""),
        })
    return filas, errores
